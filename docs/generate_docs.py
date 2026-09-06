import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

INPUT_FILE = "fullstack-dev-test-5sets-merged.md"
OUTPUT_DIR = "."

def clean_text(text):
    # Remove backticks but keep content
    text = text.replace('`', '')
    # Remove bold/italic markers
    text = text.replace('**', '').replace('__', '').replace('*', '')
    # Remove markdown links [text](url) -> text
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    # Remove extra spaces
    text = re.sub(r'  +', ' ', text)
    text = text.strip()
    return text

def set_cell_font(run, name='Times New Roman', size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold

def add_header(doc, set_num):
    # Centered title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('SOFTWARE ENGINEER ASSESSMENT')
    set_cell_font(run, bold=True, size=16)
    p.space_after = Pt(12)

    # Set info
    info = [
        f'Set: {set_num}',
        'Time: ___________ minutes',
        'Candidate Name: ____________________',
        'Date: ____________________'
    ]
    for line in info:
        p = doc.add_paragraph()
        run = p.add_run(line)
        set_cell_font(run, size=12)
        p.space_after = Pt(2)
    
    # Separator
    p = doc.add_paragraph()
    run = p.add_run('_' * 72)
    set_cell_font(run, size=10)
    p.space_after = Pt(6)

def add_section_title(doc, title, size=14):
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_cell_font(run, bold=True, size=size)
    p.space_before = Pt(12)
    p.space_after = Pt(8)

def add_mcq_question(doc, q_num, question_text, options):
    # Question
    p = doc.add_paragraph()
    run = p.add_run(f'Q{q_num}. {question_text}')
    set_cell_font(run, size=12)
    p.space_after = Pt(4)
    p.space_before = Pt(6)
    
    for opt in options:
        p = doc.add_paragraph()
        run = p.add_run(opt)
        set_cell_font(run, size=12)
        p.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(1.27)
    
    # blank line after question
    p = doc.add_paragraph()
    p.space_after = Pt(2)

def add_essay_question(doc, q_num, title, context, requirements, questions):
    # Essay title
    p = doc.add_paragraph()
    run = p.add_run(f'Q{q_num}. {title}')
    set_cell_font(run, bold=True, size=12)
    p.space_before = Pt(10)
    p.space_after = Pt(4)

    # Context
    if context:
        p = doc.add_paragraph()
        run = p.add_run(context)
        set_cell_font(run, size=12)
        p.space_after = Pt(4)

    # Requirements (Yeu cau)
    if requirements:
        p = doc.add_paragraph()
        run = p.add_run(requirements)
        set_cell_font(run, size=12)
        p.space_after = Pt(4)

    # Questions
    for q in questions:
        p = doc.add_paragraph()
        run = p.add_run(q)
        set_cell_font(run, size=12)
        p.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.64)

    # Answer space
    for _ in range(4):
        p = doc.add_paragraph()
        run = p.add_run('_' * 72)
        set_cell_font(run, size=10)

def add_answer_key(doc, answers, explanations=None):
    doc.add_page_break()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('ANSWER KEY')
    set_cell_font(run, bold=True, size=16)
    p.space_after = Pt(16)

    for item in answers:
        p = doc.add_paragraph()
        text = item
        if explanations and item[0] in explanations:
            text = f'{item} - {explanations[item[0]]}'
        run = p.add_run(text)
        set_cell_font(run, size=12)
        p.space_after = Pt(4)

def parse_markdown(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    # Find all set starting lines
    set_starts = []
    for i, line in enumerate(lines):
        m = re.match(r'^## Set (\d+)$', line.strip())
        if m:
            set_starts.append((i, int(m.group(1))))
    
    # Find answer key section
    answer_start = None
    for i, line in enumerate(lines):
        if line.strip() == '## Đáp án':
            answer_start = i
            break
    
    # Parse answers for each set - stop before Red flags / Rubric
    answers = {}
    if answer_start:
        current_set = None
        for line in lines[answer_start:]:
            stripped = line.strip()
            if stripped.startswith('## Rubric') or stripped.startswith('**Red flags'):
                break
            m = re.match(r'^### Set (\d+)$', stripped)
            if m:
                current_set = int(m.group(1))
                answers[current_set] = ''
            elif current_set and stripped and not stripped.startswith('#') and not stripped.startswith('>'):
                # Only accept line that looks like answers (has number.letter. pattern)
                if re.search(r'\d+\.[A-D]\.', stripped):
                    answers[current_set] = stripped
    
    sets_data = []
    for idx, (start_line, set_num) in enumerate(set_starts):
        end_line = set_starts[idx + 1][0] if idx + 1 < len(set_starts) else (answer_start if answer_start else len(lines))
        
        set_lines = lines[start_line:end_line]
        sets_data.append(parse_set(set_lines, set_num, answers.get(set_num, '')))
    
    return sets_data

def parse_set(lines, set_num, answer_str):
    mcqs = []
    essays = []
    
    # Parse answers
    answer_list = []
    if answer_str:
        parts = answer_str.split('.')
        i = 0
        while i < len(parts):
            part = parts[i].strip()
            m = re.match(r'^(\d+)[\s.]*([A-D])', part)
            if m:
                qnum = int(m.group(1))
                letter = m.group(2)
                answer_list.append(f'Q{qnum}. {letter}')
            elif part and re.match(r'^\d+$', part):
                # Next part starts with number
                pass
            i += 1
    
    # Fallback: parse from answer string pattern
    if not answer_list:
        # Pattern: 1.B. 2.A. 3.C. ...
        pattern = r'(\d+)\.\s*([A-D])'
        for m in re.finditer(pattern, answer_str):
            answer_list.append(f'Q{m.group(1)}. {m.group(2)}')
    
    text = '\n'.join(lines)
    
    # Split into MCQ and Essay sections
    essay_section_start = None
    for i, line in enumerate(lines):
        if re.match(r'^### Essay', line.strip()):
            essay_section_start = i
            break
    
    mcq_lines = lines[1:essay_section_start] if essay_section_start else lines[1:]
    essay_lines = lines[essay_section_start:] if essay_section_start else []
    
    # Parse MCQ questions
    current_q = None
    current_opts = []
    
    i = 0
    while i < len(mcq_lines):
        line = mcq_lines[i].strip()
        
        # Check for question start: **N.** or **N.** [tag]
        m = re.match(r'\*\*(\d+)\.\*\*\s*(.*)', line)
        if not m:
            m = re.match(r'\*\*(\d+)\.\*\*\[(.+?)\]\s*(.*)', line)
            if m:
                line = f'**{m.group(1)}.** [{m.group(2)}] {m.group(3)}'
                m = re.match(r'\*\*(\d+)\.\*\*\s*(.*)', line)
        
        if m:
            if current_q is not None:
                mcqs.append((current_q_num, current_q, current_opts))
            
            current_q_num = int(m.group(1))
            current_q = clean_text(m.group(2)).strip()
            # Remove leading tag like [debug] [trade-off] etc.
            current_q = re.sub(r'^\[.*?\]\s*', '', current_q)
            current_opts = []
            i += 1
            continue
        
        # Check for options A. B. C. D.
        opt_match = re.match(r'^([A-D])\.\s*(.*)', line)
        if opt_match and current_q is not None:
            current_opts.append(f'{opt_match.group(1)}. {clean_text(opt_match.group(2))}')
            i += 1
            continue
        
        # Continuation of question text
        if current_q is not None and line and not line.startswith('---'):
            current_q += ' ' + clean_text(line)
        
        i += 1
    
    if current_q is not None:
        mcqs.append((current_q_num, current_q, current_opts))
    
    # Parse essay sections
    essay_blocks = []
    current_essay = None
    current_context = []
    current_requirements = []
    current_questions = []
    state = 'none'  # none, context, requirements, questions
    
    for line in essay_lines:
        stripped = line.strip()
        
        if re.match(r'^### Essay (\d+)', stripped):
            if current_essay:
                essay_blocks.append({
                    'title': current_essay,
                    'context': '\n'.join(current_context).strip(),
                    'requirements': '\n'.join(current_requirements).strip(),
                    'questions': current_questions[:]
                })
            m = re.match(r'^### Essay (\d+)\s*[—\-–]\s*(.*)', stripped)
            if m:
                current_essay = clean_text(m.group(2).strip())
            else:
                current_essay = clean_text(re.sub(r'^### Essay \d+\s*[—\-–]?\s*', '', stripped))
            current_context = []
            current_requirements = []
            current_questions = []
            state = 'none'
            continue
        
        if stripped == '**Context:**':
            state = 'context'
            continue
        if stripped == '**Yêu cầu:**':
            state = 'requirements'
            continue
        if stripped == '**Câu hỏi (≤500 từ):**' or stripped.startswith('**Câu hỏi'):
            state = 'questions'
            continue
        if stripped == '---' or stripped == '':
            if state == 'none':
                continue
        
        if state == 'context':
            current_context.append(clean_text(stripped))
        elif state == 'requirements':
            if stripped.startswith('- **') or stripped.startswith('-'):
                current_requirements.append(clean_text(stripped.lstrip('- ')))
            else:
                current_requirements.append(clean_text(stripped))
        elif state == 'questions':
            if stripped and (stripped[0].isdigit() and '. ' in stripped[:4]):
                current_questions.append(clean_text(stripped))
            elif stripped:
                # continuation
                if current_questions:
                    current_questions[-1] += ' ' + clean_text(stripped)
                else:
                    current_questions.append(clean_text(stripped))
    
    if current_essay:
        essay_blocks.append({
            'title': current_essay,
            'context': '\n'.join(current_context).strip(),
            'requirements': '\n'.join(current_requirements).strip(),
            'questions': current_questions[:]
        })
    
    return {
        'set_num': set_num,
        'mcqs': mcqs,
        'essays': essay_blocks,
        'answers': answer_list
    }

def generate_docx(set_data):
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    
    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    
    # Page 1: Header
    add_header(doc, set_data['set_num'])
    
    # Section 1: Multiple Choice
    doc.add_page_break()
    add_section_title(doc, 'Section 1: Multiple Choice')
    
    for q_num, q_text, options in set_data['mcqs']:
        add_mcq_question(doc, q_num, q_text, options)
    
    # Section 2: Essay
    if set_data['essays']:
        doc.add_page_break()
        add_section_title(doc, 'Section 2: Essay')
        
        for idx, essay in enumerate(set_data['essays']):
            ctx = essay.get('context', '')
            req = essay.get('requirements', '')
            qs = essay.get('questions', [])
            add_essay_question(doc, idx + 1, essay['title'], ctx, req, qs)
    
    # Final page: Answer Key
    add_answer_key(doc, set_data['answers'])
    
    # Save
    filename = f'set_{set_data["set_num"]}.docx'
    filepath = os.path.join(OUTPUT_DIR, filename)
    doc.save(filepath)
    print(f'Created: {filepath}')
    return filepath

def main():
    with open(INPUT_FILE, 'r', encoding='utf-8') as f:
        content = f.read()
    
    sets_data = parse_markdown(INPUT_FILE)
    
    print(f'Found {len(sets_data)} sets')
    for sd in sets_data:
        print(f'  Set {sd["set_num"]}: {len(sd["mcqs"])} MCQs, {len(sd["essays"])} essays, {len(sd["answers"])} answers')
        filepath = generate_docx(sd)
        print(f'    -> {filepath}')

if __name__ == '__main__':
    main()
